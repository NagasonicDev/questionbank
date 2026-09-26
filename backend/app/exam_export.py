"""
Test generator — turns a set of questions into a downloadable DOCX or PDF
that looks like an actual exam paper: title block, instructions, numbered
questions (unlike the in-app reader, a printed exam is expected to number
its questions), equations rendered as real typeset math, embedded images/
diagrams, tables, and an optional answer/solution section at the end.

Equations are rendered via matplotlib's mathtext engine (no LaTeX
installation required) to a transparent PNG and embedded as an image —
this covers the common exam math (fractions, powers, roots, sums,
integrals, Greek letters, sub/superscripts) without needing a system LaTeX.
Anything mathtext can't parse falls back to plain monospaced text rather
than failing the whole export.
"""
import io
import random
import re
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle,
    PageBreak, HRFlowable, ListFlowable, ListItem, KeepTogether,
)

from .database import ASSETS_DIR
from . import models

GRAY = RGBColor(0x6B, 0x6B, 0x64)


# ===================== Equation rendering (shared) =====================

def render_latex_png(latex: str, fontsize: int = 15) -> bytes | None:
    text = (latex or "").strip()
    if not text:
        return None
    if not (text.startswith("$") and text.endswith("$")):
        text = f"${text}$"
    fig = plt.figure(figsize=(0.01, 0.01))
    try:
        fig.text(0, 0, text, fontsize=fontsize, color="black")
        fig.canvas.draw()
        renderer = fig.canvas.get_renderer()
        t = fig.texts[0]
        bbox = t.get_window_extent(renderer=renderer)
        width_in, height_in = bbox.width / fig.dpi, bbox.height / fig.dpi
        fig.set_size_inches(max(width_in, 0.1) + 0.06, max(height_in, 0.1) + 0.06)
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=220, transparent=True, bbox_inches="tight", pad_inches=0.03)
        plt.close(fig)
        buf.seek(0)
        return buf.read()
    except Exception:
        plt.close(fig)
        return None


def _asset_bytes(asset_path: str) -> bytes | None:
    p = ASSETS_DIR / asset_path
    if p.exists() and p.is_file():
        return p.read_bytes()
    return None


# ===================== Question selection helper (used by the router) =====================

def blocks_by_slot(question: models.Question, slot: str) -> list[models.ContentBlock]:
    return sorted([b for b in question.content_blocks if b.slot == slot], key=lambda b: b.position)


def _generated_part_answer_lines(question: models.Question, part: models.Question) -> int:
    if question.type_key not in {"extended_response", "short_response", "short_answer"}:
        return 0
    blocks = blocks_by_slot(part, "body")
    if any(block.block_type == "answer_area" for block in blocks):
        return 0
    return max(1, int((part.marks or 0) * 2 + 0.999999))


def select_questions_for_marks(
    db, *, course_id: str, node_ids: list[str] | None, type_key: str | None,
    difficulty_min: int | None, difficulty_max: int | None, target_marks: float, shuffle: bool = True,
    marks_min: float | None = None, marks_max: float | None = None,
    type_keys: list[str] | None = None, difficulties: list[int] | None = None,
) -> tuple[list[models.Question], float]:
    """Pick a set of approved, whole (non-part) questions whose marks add up
    to the closest reachable total to `target_marks`, preferring not to
    overshoot when two totals are equally close.

    Only questions with a marks value are eligible, since there's nothing to
    sum for a question with no marks recorded.
    """
    from . import crud  # local import: avoids a circular import at module load time

    stmt = crud.apply_question_filters(
        select(models.Question.question_id, models.Question.marks, models.Question.created_at), course_id=course_id, node_ids=node_ids, type_key=type_key,
        difficulty_min=difficulty_min, difficulty_max=difficulty_max, tag_name=None,
        marks_min=marks_min, marks_max=marks_max,
        type_keys=type_keys, difficulties=difficulties,
    ).where(models.Question.marks.is_not(None))

    candidates = list(db.execute(stmt).all())
    if shuffle:
        random.shuffle(candidates)
    else:
        candidates.sort(key=lambda row: row[2])

    target_units = max(1, round(target_marks * 100))
    questions_by_mark: dict[int, list[str]] = {}
    for question_id, marks, _created_at in candidates:
        units = max(1, round(marks * 100))
        questions_by_mark.setdefault(units, []).append(question_id)

    bundles: list[tuple[int, list[str]]] = []
    for units, questions in questions_by_mark.items():
        offset = 0
        power = 1
        while offset < len(questions):
            count = min(power, len(questions) - offset)
            bundles.append((units * count, questions[offset:offset + count]))
            offset += count
            power *= 2

    # Questions above the target only need singleton consideration: adding
    # anything else makes their overshoot worse. Keep the search bounded at 2×target.
    max_question_units = max((units for units in questions_by_mark if units <= target_units), default=0)
    limit = target_units + max_question_units
    reachable = [0]
    seen = {0}
    parents: dict[int, tuple[int, int]] = {}
    best = 0
    best_distance: int | None = None
    best_path: list[int] = []
    for bundle_index, (bundle_marks, bundle_questions) in enumerate(bundles):
        if len(bundle_questions) == 1:
            distance = abs(bundle_marks - target_units)
            if best_distance is None or distance < best_distance or (distance == best_distance and bundle_marks < best):
                best = bundle_marks
                best_distance = distance
                best_path = [bundle_index]

    for bundle_index, (bundle_marks, _) in enumerate(bundles):
        for previous in reachable.copy():
            total = previous + bundle_marks
            if total > limit or total in seen:
                continue
            seen.add(total)
            reachable.append(total)
            parents[total] = (previous, bundle_index)
            distance = abs(total - target_units)
            if best_distance is None or distance < best_distance or (distance == best_distance and total < best):
                best = total
                best_distance = distance
                best_path = [bundle_index]
                cursor = previous
                while cursor > 0:
                    cursor, parent_bundle = parents[cursor]
                    best_path.append(parent_bundle)
            if best_distance == 0:
                break
        if best_distance == 0:
            break

    selected_ids = [question_id for bundle_index in best_path for question_id in bundles[bundle_index][1]]
    if selected_ids:
        selected_by_id: dict[str, models.Question] = {}
        for offset in range(0, len(selected_ids), 500):
            batch = selected_ids[offset:offset + 500]
            selected_by_id.update({
                question.question_id: question
                for question in db.scalars(
                    crud.hydrated_question_query().where(models.Question.question_id.in_(batch))
                ).unique().all()
            })
        selected = [selected_by_id[question_id] for question_id in selected_ids]
    else:
        selected = []

    if shuffle:
        random.shuffle(selected)
    return selected, sum(q.marks or 0 for q in selected)


def select_n_questions(
    db, *, course_id: str, node_ids: list[str] | None, type_key: str | None,
    difficulty_min: int | None, difficulty_max: int | None, count: int, shuffle: bool = True,
    marks_min: float | None = None, marks_max: float | None = None,
    type_keys: list[str] | None = None, difficulties: list[int] | None = None,
) -> list[models.Question]:
    """Pick up to `count` whole, approved questions matching the filters.

    Unlike marks-based selection, question marks are irrelevant here — this
    is for sections like "15 multiple choice questions" where the requirement
    is a count of questions of a given type, not a mark total. If fewer than
    `count` eligible questions exist, returns what's available.
    """
    from . import crud  # local import: avoids a circular import at module load time

    stmt = crud.apply_question_filters(
        crud.hydrated_question_query(), course_id=course_id, node_ids=node_ids, type_key=type_key,
        difficulty_min=difficulty_min, difficulty_max=difficulty_max, tag_name=None,
        marks_min=marks_min, marks_max=marks_max,
        type_keys=type_keys, difficulties=difficulties,
    )

    candidates = list(db.scalars(stmt).unique().all())
    if shuffle:
        random.shuffle(candidates)
    else:
        candidates.sort(key=lambda q: q.created_at)
    return candidates[:count]


# ===================== Marking-criteria table =====================
#
# Marking guides are stored as free-form blocks (usually list/text). When
# rendered — in-app, in the DOCX solutions doc, or in the PDF — they're
# displayed as a two-column table: criteria on the left, the mark allocation
# for that criterion/band on the right, one row per mark range. Allocation is
# parsed from the conventional "N marks: ..." (or "... (N marks)") phrasing
# used by the course skill's import guidance; criteria without a parseable
# allocation just get an empty marks cell.

_MARKS_PREFIX_RE = re.compile(
    r"^\s*(\d+(?:\.\d+)?)\s*"
    r"(?:(?:-\s*|–\s*|—\s*|\bto\b\s+|\bor\b\s+|\/\s*)(\d+(?:\.\d+)?))?"
    r"\s*marks?\b",
    re.IGNORECASE,
)
_MARKS_TRAIL_RE = re.compile(
    r"\s*[\(\[]([\d\s.\-–—/]+)\s*marks?[\)\]]\s*$",
    re.IGNORECASE,
)


def _extract_marks(text: str) -> tuple[str, str]:
    """Split a criterion line into (mark allocation, clean criteria text).

    Handles leading allocations ("3 marks: ...", "1–2 marks for ...") and
    trailing ones ("Correct substitution (2 marks)").
    """
    t = text.strip()
    m = _MARKS_PREFIX_RE.match(t)
    if m:
        low, high = m.group(1), m.group(2)
        alloc = f"{low}\u2013{high}" if high else low
        rest = t[m.end():].lstrip().lstrip(":–—:-").strip()
        return alloc, rest
    m = _MARKS_TRAIL_RE.search(t)
    if m:
        return m.group(1).strip().replace(" ", ""), t[:m.start()].strip()
    return "", t


def criteria_rows(blocks: list[models.ContentBlock]) -> list[dict]:
    """Convert marking-criteria content blocks into table rows.

    list/text blocks become one row per criterion; any other block type
    (equations, images, tables, ...) produces no rows and is rendered
    separately by the caller so nothing is ever lost.
    """
    import json as _json

    rows: list[dict] = []
    for b in blocks:
        content = _json.loads(b.content_json)
        if b.block_type == "list":
            for item in content.get("items", []):
                alloc, crit = _extract_marks(str(item))
                rows.append({"criteria": crit if crit else str(item).strip(), "marks": alloc})
        elif b.block_type == "text":
            text = (content.get("text", "") or "").strip()
            if text:
                alloc, crit = _extract_marks(text)
                rows.append({"criteria": crit if crit else text, "marks": alloc})
    return rows


def _criteria_supplementary_blocks(blocks: list[models.ContentBlock]) -> list[models.ContentBlock]:
    """Blocks inside a marking guide that don't fit the criteria table
    (equations, images, tables, code) — rendered after the table."""
    return [b for b in blocks if b.block_type not in ("list", "text")]


# =====================================================================
# DOCX
# =====================================================================

def _docx_add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    borders.append(bottom)
    p_pr.append(borders)


def _docx_render_blocks(doc: Document, blocks: list[models.ContentBlock], indent: float = 0.0):
    import json as _json
    for block in blocks:
        content = _json.loads(block.content_json)
        bt = block.block_type

        if bt == "text":
            p = doc.add_paragraph(content.get("text", ""))
            p.paragraph_format.left_indent = Inches(indent)

        elif bt == "heading":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(indent)
            run = p.add_run(content.get("text", ""))
            run.bold = True
            run.font.size = Pt(13)

        elif bt == "equation":
            png = render_latex_png(content.get("latex", ""))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(indent)
            if content.get("display", True):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if png:
                run = p.add_run()
                run.add_picture(io.BytesIO(png), height=Pt(20))
            else:
                run = p.add_run(content.get("latex", ""))
                run.font.name = "Courier New"

        elif bt in ("image", "diagram", "graph"):
            data = _asset_bytes(content.get("asset_path", ""))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(indent)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if data:
                run = p.add_run()
                run.add_picture(io.BytesIO(data), width=Inches(min(4.5, 6.5 - indent)))
            else:
                run = p.add_run(f"[missing image: {content.get('asset_path', '?')}]")
                run.italic = True
            if content.get("caption"):
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cr = cap.add_run(content["caption"])
                cr.italic = True
                cr.font.size = Pt(9)
                cr.font.color.rgb = GRAY

        elif bt == "table":
            columns = content.get("columns", [])
            rows = content.get("rows", [])
            if columns:
                table = doc.add_table(rows=1 + len(rows), cols=len(columns))
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for j, col in enumerate(columns):
                    cell = table.rows[0].cells[j]
                    cell.text = str(col)
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True
                for i, row in enumerate(rows):
                    for j, val in enumerate(row):
                        if j < len(columns):
                            table.rows[i + 1].cells[j].text = str(val)

        elif bt == "list":
            style = "List Number" if content.get("ordered") else "List Bullet"
            for item in content.get("items", []):
                p = doc.add_paragraph(item, style=style)
                p.paragraph_format.left_indent = Inches(indent + 0.25)

        elif bt == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(indent + 0.15)
            run = p.add_run(content.get("code", ""))
            run.font.name = "Courier New"
            run.font.size = Pt(10)

        elif bt == "answer_area":
            for _ in range(int(content.get("lines", 3))):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(indent)
                p.paragraph_format.space_after = Pt(14)
                _docx_add_bottom_border(p)

        elif bt == "page_break":
            doc.add_page_break()


def _docx_render_criteria_table(doc: Document, blocks: list[models.ContentBlock]):
    rows = criteria_rows(blocks)
    if not rows:
        return
    widths = (Inches(4.6), Inches(0.9))
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(("Criteria", "Marks")):
        cell = table.rows[0].cells[j]
        cell.text = h
        cell.width = widths[j]
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
    for i, row in enumerate(rows):
        cell = table.rows[i + 1].cells[0]
        cell.text = row["criteria"]
        cell.width = widths[0]
        cell = table.rows[i + 1].cells[1]
        cell.text = row["marks"]
        cell.width = widths[1]
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _docx_render_question(doc: Document, number: int, q: models.Question):
    header = doc.add_paragraph()
    header.paragraph_format.space_before = Pt(18)
    hr = header.add_run(f"Question {number}")
    hr.bold = True
    hr.font.size = Pt(12)
    if q.marks is not None:
        mr = header.add_run(f"\t[{q.marks:g} mark{'s' if q.marks != 1 else ''}]")
        mr.font.size = Pt(10)
        mr.font.color.rgb = GRAY
        header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), alignment=WD_TAB_ALIGNMENT.RIGHT)

    question_body = blocks_by_slot(q, "body")
    if q.children:
        question_body = [block for block in question_body if block.block_type != "answer_area"]
    _docx_render_blocks(doc, question_body)

    for part in q.children:
        part_p = doc.add_paragraph()
        part_p.paragraph_format.left_indent = Inches(0.3)
        pr = part_p.add_run(f"({part.part_label})")
        pr.bold = True
        if part.marks is not None:
            part_p.add_run(f"  [{part.marks:g} mark{'s' if part.marks != 1 else ''}]").font.size = Pt(9)
        _docx_render_blocks(doc, blocks_by_slot(part, "body"), indent=0.3)
        for _ in range(_generated_part_answer_lines(q, part)):
            answer_line = doc.add_paragraph()
            answer_line.paragraph_format.left_indent = Inches(0.3)
            answer_line.paragraph_format.space_after = Pt(14)
            _docx_add_bottom_border(answer_line)

    if q.source:
        src_p = doc.add_paragraph()
        parts_txt = [q.source.name] + ([str(q.source.year)] if q.source.year else [])
        sr2 = src_p.add_run(f"(Source: {', '.join(parts_txt)})")
        sr2.italic = True
        sr2.font.size = Pt(8)
        sr2.font.color.rgb = GRAY


def _docx_render_section_heading(doc: Document, name: str, questions: list[models.Question]):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run(name)
    r.bold = True
    r.font.size = Pt(14)
    sub = doc.add_paragraph()
    sec_marks = sum(q.marks or 0 for q in questions)
    subr = sub.add_run(f"{len(questions)} question{'s' if len(questions) != 1 else ''}  ·  {sec_marks:g} marks")
    subr.font.size = Pt(10)
    subr.font.color.rgb = GRAY
    _docx_add_bottom_border(sub)


def build_docx(*, title: str, course_name: str, questions: list[models.Question],
               include_answers: bool, sections: list[tuple[str | None, list[models.Question]]] | None = None) -> bytes:
    if sections is None:
        sections = [(None, questions)]
    all_questions = [q for _, qs in sections for q in qs]

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for m in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin):
        pass
    section.top_margin = section.bottom_margin = Inches(0.9)
    section.left_margin = section.right_margin = Inches(0.9)

    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title_p.add_run(title)
    tr.bold = True
    tr.font.size = Pt(20)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub_p.add_run(f"{course_name}  ·  {date.today().strftime('%d %B %Y')}  ·  {len(all_questions)} questions")
    sr.font.size = Pt(10)
    sr.font.color.rgb = GRAY

    total_marks = sum(q.marks or 0 for q in all_questions)
    if total_marks:
        marks_p = doc.add_paragraph()
        marks_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mr = marks_p.add_run(f"Total marks: {total_marks:g}")
        mr.font.size = Pt(10)
        mr.font.color.rgb = GRAY

    doc.add_paragraph()
    instr = doc.add_paragraph()
    ir = instr.add_run("Answer all questions in the spaces provided. Show all working where relevant.")
    ir.italic = True
    ir.font.size = Pt(10)

    rule_p = doc.add_paragraph()
    _docx_add_bottom_border(rule_p)

    number = 0
    for name, qs in sections:
        if name:
            _docx_render_section_heading(doc, name, qs)
        for q in qs:
            number += 1
            _docx_render_question(doc, number, q)

    if include_answers:
        doc.add_page_break()
        _docx_render_solutions_section(doc, all_questions, heading="Answers", sections=sections)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_render_solutions_section(doc: Document, questions: list[models.Question],
                                   heading: str | None = "Solutions & Marking Guide",
                                   sections: list[tuple[str | None, list[models.Question]]] | None = None):
    if heading:
        ans_title = doc.add_paragraph()
        ans_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        atr = ans_title.add_run(heading)
        atr.bold = True
        atr.font.size = Pt(18)
        doc.add_paragraph()

    if sections is None:
        sections = [(None, questions)]

    number = 0
    for name, sec_questions in sections:
        if name:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            r = p.add_run(name)
            r.bold = True
            r.font.size = Pt(13)
        for q in sec_questions:
            number += 1
            header = doc.add_paragraph()
            header.paragraph_format.space_before = Pt(14)
            hr = header.add_run(f"Question {number}")
            hr.bold = True
            hr.font.size = Pt(12)
            if q.marks is not None:
                mr = header.add_run(f"\t[{q.marks:g} mark{'s' if q.marks != 1 else ''}]")
                mr.font.size = Pt(10)
                mr.font.color.rgb = GRAY
                header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), alignment=WD_TAB_ALIGNMENT.RIGHT)

            marking_blocks = blocks_by_slot(q, "marking_criteria")
            answer_blocks = blocks_by_slot(q, "answer")
            solution_blocks = blocks_by_slot(q, "solution")

            # Marking guide comes first and is the primary grading reference —
            # it's what a marker actually checks a response against. It's
            # rendered as a criteria/marks table when the content is list/text.
            if marking_blocks:
                mg = doc.add_paragraph()
                mg.add_run("Marking guide:").italic = True
                _docx_render_criteria_table(doc, marking_blocks)
                supplementary = _criteria_supplementary_blocks(marking_blocks)
                if supplementary:
                    doc.add_paragraph().add_run("Detailed marking notes:").italic = True
                    _docx_render_blocks(doc, supplementary, indent=0.2)
            if answer_blocks:
                doc.add_paragraph().add_run("Answer:").italic = True
                _docx_render_blocks(doc, answer_blocks, indent=0.2)
            if solution_blocks:
                doc.add_paragraph().add_run("Solution / working:").italic = True
                _docx_render_blocks(doc, solution_blocks, indent=0.2)
            if not marking_blocks and not answer_blocks and not solution_blocks:
                np = doc.add_paragraph()
                np.add_run("No marking guide recorded for this question.").italic = True


def build_solutions_docx(*, title: str, course_name: str, questions: list[models.Question],
                         sections: list[tuple[str | None, list[models.Question]]] | None = None) -> bytes:
    """Standalone marking-guide document: no question bodies, just each
    question's marking criteria / answer / solution, numbered to match the
    corresponding test paper."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.9)
    section.left_margin = section.right_margin = Inches(0.9)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title_p.add_run(f"Solutions — {title}")
    tr.bold = True
    tr.font.size = Pt(18)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub_p.add_run(f"{course_name}  ·  {date.today().strftime('%d %B %Y')}")
    sr.font.size = Pt(10)
    sr.font.color.rgb = GRAY
    doc.add_paragraph()

    _docx_render_solutions_section(doc, questions, heading=None, sections=sections)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# =====================================================================
# PDF (reportlab)
# =====================================================================

def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("ExamTitle", parent=styles["Title"], fontSize=20, alignment=TA_CENTER))
    styles.add(ParagraphStyle("ExamSubtitle", parent=styles["Normal"], fontSize=10, alignment=TA_CENTER, textColor=colors.HexColor("#6b6b64")))
    styles.add(ParagraphStyle("Instructions", parent=styles["Normal"], fontSize=10, italic=True, spaceAfter=10))
    styles.add(ParagraphStyle("QHeader", parent=styles["Normal"], fontSize=12, spaceBefore=16, spaceAfter=6, fontName="Helvetica-Bold"))
    styles.add(ParagraphStyle("Body", parent=styles["Normal"], fontSize=11, spaceAfter=6, alignment=TA_JUSTIFY))
    styles.add(ParagraphStyle("BodyIndent", parent=styles["Body"], leftIndent=18))
    styles.add(ParagraphStyle("SourceLine", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#6b6b64")))
    styles.add(ParagraphStyle("Caption", parent=styles["Normal"], fontSize=9, alignment=TA_CENTER, textColor=colors.HexColor("#6b6b64")))
    styles.add(ParagraphStyle("ExamCode", parent=styles["Normal"], fontName="Courier", fontSize=9, leftIndent=12))
    return styles


def _pdf_render_blocks(story: list, styles, blocks: list[models.ContentBlock], indent: bool = False):
    import json as _json
    body_style = styles["BodyIndent"] if indent else styles["Body"]

    for block in blocks:
        content = _json.loads(block.content_json)
        bt = block.block_type

        if bt == "text":
            story.append(Paragraph(content.get("text", ""), body_style))

        elif bt == "heading":
            story.append(Paragraph(f"<b>{content.get('text', '')}</b>", body_style))

        elif bt == "equation":
            png = render_latex_png(content.get("latex", ""))
            if png:
                img = RLImage(io.BytesIO(png))
                img._restrictSize(3 * inch, 0.32 * inch)
                if content.get("display", True):
                    img.hAlign = "CENTER"
                story.append(img)
            else:
                story.append(Paragraph(f"<font face='Courier'>{content.get('latex', '')}</font>", body_style))

        elif bt in ("image", "diagram", "graph"):
            data = _asset_bytes(content.get("asset_path", ""))
            if data:
                img = RLImage(io.BytesIO(data))
                img._restrictSize(4.5 * inch, 4.5 * inch)
                img.hAlign = "CENTER"
                story.append(img)
                if content.get("caption"):
                    story.append(Paragraph(content["caption"], styles["Caption"]))
            else:
                story.append(Paragraph(f"<i>[missing image: {content.get('asset_path', '?')}]</i>", body_style))

        elif bt == "table":
            columns = content.get("columns", [])
            rows = content.get("rows", [])
            if columns:
                data = [columns] + rows
                t = Table(data, hAlign="CENTER")
                t.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f1ee")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                ]))
                story.append(t)
                story.append(Spacer(1, 6))

        elif bt == "list":
            items = [ListItem(Paragraph(item, body_style)) for item in content.get("items", [])]
            story.append(ListFlowable(items, bulletType="1" if content.get("ordered") else "bullet"))

        elif bt == "code":
            for line in content.get("code", "").split("\n"):
                story.append(Paragraph(line.replace(" ", "&nbsp;") or "&nbsp;", styles["ExamCode"]))

        elif bt == "answer_area":
            for _ in range(int(content.get("lines", 3))):
                t = Table([[""]], colWidths=[6.3 * inch], rowHeights=[0.35 * inch])
                t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.75, colors.HexColor("#999999"))]))
                story.append(t)

        elif bt == "page_break":
            story.append(PageBreak())


def _pdf_criteria_table(blocks: list[models.ContentBlock]):
    rows = criteria_rows(blocks)
    if not rows:
        return None
    data = [["Criteria", "Marks"]] + [[r["criteria"], r["marks"]] for r in rows]
    t = Table(data, colWidths=[5.0 * inch, 1.0 * inch], hAlign="LEFT")
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f1ee")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (1, 0), (1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def _pdf_render_question(story: list, styles, number: int, q: models.Question):
    marks_txt = f"[{q.marks:g} mark{'s' if q.marks != 1 else ''}]" if q.marks is not None else ""
    header_tbl = Table([[Paragraph(f"<b>Question {number}</b>", styles["QHeader"]),
                          Paragraph(marks_txt, styles["SourceLine"])]],
                        colWidths=[5.3 * inch, 1 * inch])
    header_tbl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("ALIGN", (1, 0), (1, 0), "RIGHT")]))

    block_story: list = [header_tbl]
    question_body = blocks_by_slot(q, "body")
    if q.children:
        question_body = [block for block in question_body if block.block_type != "answer_area"]
    _pdf_render_blocks(block_story, styles, question_body)

    for part in q.children:
        part_marks = f"  [{part.marks:g} mark{'s' if part.marks != 1 else ''}]" if part.marks is not None else ""
        block_story.append(Paragraph(f"<b>({part.part_label})</b>{part_marks}", styles["BodyIndent"]))
        _pdf_render_blocks(block_story, styles, blocks_by_slot(part, "body"), indent=True)
        for _ in range(_generated_part_answer_lines(q, part)):
            answer_line = Table([[""]], colWidths=[6.0 * inch], rowHeights=[0.35 * inch])
            answer_line.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.75, colors.HexColor("#999999"))]))
            block_story.append(answer_line)

    if q.source:
        parts_txt = [q.source.name] + ([str(q.source.year)] if q.source.year else [])
        block_story.append(Paragraph(f"(Source: {', '.join(parts_txt)})", styles["SourceLine"]))

    story.append(KeepTogether(block_story[:2]) if len(block_story) > 1 else block_story[0])
    story.extend(block_story[2:] if len(block_story) > 2 else [])


def _pdf_render_section_heading(story: list, styles, name: str, questions: list[models.Question]):
    sec_marks = sum(q.marks or 0 for q in questions)
    story.append(Paragraph(
        f"<b>{name}</b>",
        ParagraphStyle("SectionTitle", parent=styles["QHeader"], fontSize=14, spaceBefore=20, spaceAfter=2),
    ))
    story.append(Paragraph(
        f"{len(questions)} question{'s' if len(questions) != 1 else ''} &middot; {sec_marks:g} marks",
        styles["SourceLine"],
    ))
    story.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc")))
    story.append(Spacer(1, 4))


def build_pdf(*, title: str, course_name: str, questions: list[models.Question], include_answers: bool,
              sections: list[tuple[str | None, list[models.Question]]] | None = None) -> bytes:
    if sections is None:
        sections = [(None, questions)]
    all_questions = [q for _, qs in sections for q in qs]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=LETTER,
        topMargin=0.9 * inch, bottomMargin=0.9 * inch, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
    )
    styles = _pdf_styles()
    story: list = []

    story.append(Paragraph(title, styles["ExamTitle"]))
    total_marks = sum(q.marks or 0 for q in all_questions)
    subtitle = f"{course_name} &middot; {date.today().strftime('%d %B %Y')} &middot; {len(all_questions)} questions"
    if total_marks:
        subtitle += f" &middot; Total marks: {total_marks:g}"
    story.append(Paragraph(subtitle, styles["ExamSubtitle"]))
    story.append(Spacer(1, 10))
    story.append(Paragraph("Answer all questions in the spaces provided. Show all working where relevant.", styles["Instructions"]))
    story.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc")))

    number = 0
    for name, qs in sections:
        if name:
            _pdf_render_section_heading(story, styles, name, qs)
        for q in qs:
            number += 1
            _pdf_render_question(story, styles, number, q)

    if include_answers:
        story.append(PageBreak())
        _pdf_append_solutions_section(story, styles, questions, heading="Answers", sections=sections)

    doc.build(story)
    return buf.getvalue()


def _pdf_append_solutions_section(story: list, styles, questions: list[models.Question],
                                  heading: str | None = "Solutions & Marking Guide",
                                  sections: list[tuple[str | None, list[models.Question]]] | None = None):
    if heading:
        story.append(Paragraph(heading, styles["ExamTitle"]))
        story.append(Spacer(1, 12))

    if sections is None:
        sections = [(None, questions)]

    number = 0
    for name, sec_questions in sections:
        if name:
            story.append(Paragraph(
                f"<b>{name}</b>",
                ParagraphStyle("SectionTitle", parent=styles["QHeader"], fontSize=13, spaceBefore=14, spaceAfter=2),
            ))
            story.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc")))
            story.append(Spacer(1, 4))
        for q in sec_questions:
            number += 1
            marks_txt = f"[{q.marks:g} mark{'s' if q.marks != 1 else ''}]" if q.marks is not None else ""
            header_tbl = Table([[Paragraph(f"<b>Question {number}</b>", styles["QHeader"]),
                                  Paragraph(marks_txt, styles["SourceLine"])]],
                                colWidths=[5.3 * inch, 1 * inch])
            header_tbl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("ALIGN", (1, 0), (1, 0), "RIGHT")]))
            story.append(header_tbl)

            marking_blocks = blocks_by_slot(q, "marking_criteria")
            answer_blocks = blocks_by_slot(q, "answer")
            solution_blocks = blocks_by_slot(q, "solution")

            # Marking guide first, rendered as a criteria/marks table when the
            # content is list/text; anything that can't live in the table
            # (equations, images, tables, code) follows under its own label.
            if marking_blocks:
                story.append(Paragraph("<i>Marking guide:</i>", styles["Body"]))
                table = _pdf_criteria_table(marking_blocks)
                if table is not None:
                    story.append(table)
                    story.append(Spacer(1, 4))
                supplementary = _criteria_supplementary_blocks(marking_blocks)
                if supplementary:
                    story.append(Paragraph("<i>Detailed marking notes:</i>", styles["Body"]))
                    _pdf_render_blocks(story, styles, supplementary, indent=True)
            if answer_blocks:
                story.append(Paragraph("<i>Answer:</i>", styles["Body"]))
                _pdf_render_blocks(story, styles, answer_blocks, indent=True)
            if solution_blocks:
                story.append(Paragraph("<i>Solution / working:</i>", styles["Body"]))
                _pdf_render_blocks(story, styles, solution_blocks, indent=True)
            if not marking_blocks and not answer_blocks and not solution_blocks:
                story.append(Paragraph("<i>No marking guide recorded for this question.</i>", styles["Body"]))


def build_solutions_pdf(*, title: str, course_name: str, questions: list[models.Question],
                        sections: list[tuple[str | None, list[models.Question]]] | None = None) -> bytes:
    """Standalone marking-guide PDF, matching build_solutions_docx."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=LETTER,
        topMargin=0.9 * inch, bottomMargin=0.9 * inch, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
    )
    styles = _pdf_styles()
    story: list = []
    story.append(Paragraph(f"Solutions — {title}", styles["ExamTitle"]))
    story.append(Paragraph(f"{course_name} &middot; {date.today().strftime('%d %B %Y')}", styles["ExamSubtitle"]))
    story.append(Spacer(1, 12))
    _pdf_append_solutions_section(story, styles, questions, heading=None, sections=sections)
    doc.build(story)
    return buf.getvalue()
